"""Regenerate esyluban/templates/DataTables, the sample tables every release package ships.

They show the recommended way to organise a game that ships in several languages and
switches between them at runtime:

  items.xlsx, monsters.xlsx, npcs.xlsx
                        player-visible text is a key, typed string#ref=text.TbText
  text/texts.xlsx       the text table, an ordinary table: one row per key, one column
                        per language, exported to the game as it is
  variant_en/npcs.xlsx  what the English version needs besides translation: 3001 has an
                        English voice recording, 3002 does not yet and keeps the default

docs/writing-tables.md walks through exactly these rows, so change both together.
Running it again only rewrites spreadsheets whose cells changed.

    python esyluban/scripts/authoring/create_template_tables.py
"""
from pathlib import Path

import openpyxl

DATA = Path(__file__).resolve().parents[2] / "templates" / "DataTables"
TEXT = "string#ref=text.TbText"


def cells(wb):
    # an empty string is written as a blank cell and reads back as None
    return [(ws.title, [[None if v == "" else v for v in row] for row in ws.iter_rows(values_only=True)])
            for ws in wb.worksheets]


def book(relative, b1, header, types, rows):
    path = DATA / relative
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "##export"
    ws["B1"] = b1
    ws.append(["##var"] + header)
    ws.append(["##type"] + types)
    for row in rows:
        ws.append([None] + row)
    if path.exists() and cells(openpyxl.load_workbook(path)) == cells(wb):
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def b1(full_name):
    return f'full_name="{full_name}" & read_schema_from_file="true"'


book("items.xlsx", b1("demo.TbItem"), ["id", "name", "price", "desc"], ["int", TEXT, "int", TEXT], [
    [1001, "/item_1001_name", 50, "/item_1001_desc"],
    [1002, "/item_1002_name", 300, "/item_1002_desc"],
    [1003, "/item_1003_name", 20, "/item_1003_desc"],
])

book("monsters.xlsx", b1("demo.TbMonster"), ["id", "name", "hp", "atk"], ["int", TEXT, "int", "int"], [
    [2001, "/monster_2001_name", 30, 5],
    [2002, "/monster_2002_name", 80, 12],
    [2003, "/monster_2003_name", 250, 30],
])

NPC = b1("demo.TbNpc"), ["id", "name", "voice"], ["int", TEXT, "string"]
book("npcs.xlsx", *NPC, [
    [3001, "/npc_3001_name", "Voice/zh/chief.wav"],
    [3002, "/npc_3002_name", "Voice/zh/smith.wav"],
])
book("variant_en/npcs.xlsx", *NPC, [
    [3001, "/npc_3001_name", "Voice/en/chief.wav"],
])

book("text/texts.xlsx", b1("text.TbText"), ["key", "zh", "en"], ["string", "string", "string"], [
    ["/item_1001_name", "木剑", "Wooden Sword"],
    ["/item_1001_desc", "新手用的练习剑", "A practice sword for beginners"],
    ["/item_1002_name", "铁剑", "Iron Sword"],
    ["/item_1002_desc", "标准的铁制长剑", "A standard iron sword"],
    ["/item_1003_name", "生命药水", "Health Potion"],
    ["/item_1003_desc", "恢复少量生命", "Restores a little health"],
    ["/monster_2001_name", "史莱姆", "Slime"],
    ["/monster_2002_name", "哥布林", "Goblin"],
    ["/monster_2003_name", "恶魔", "Demon"],
    ["/npc_3001_name", "村长", "Village Chief"],
    ["/npc_3002_name", "铁匠", "Blacksmith"],
])

print("ok")
